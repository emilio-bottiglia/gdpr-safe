#v1.11 see CHANGELOG.MD

from pathlib import Path
import fnmatch
import csv
from time import sleep
import openpyxl
import docx
from tqdm import tqdm
import pymupdf

#global variables to hold filenames
txt_files = []
csv_files = []
pdf_files = []
docx_files = []
excel_files = []

#gdpr sample keywords
words = ["name", "surname", "first name", "second name", "maiden name", "address",
         "passport", "driving license", "national id", "pps", "ppsn",
         "social security number", "sex", "sexual", "phone", "health", "religion",
         "photo", "genetic", "political", "straight", "gay", "bisex", "transex",
         "orientation", "location", "ip address", "child", "kid", "coordinates",
         "physical", "mental", "lgbt", "racial", "religious", "philosophical",
        "trade union", "children", "biological", "biomedical", "medical", "diagnostic",
        "physiology", "behavior", "behavioural", "facial", "visual", "dactyloscopic",
        "fingerprint", "disease", "disability", "risk", "clinical", "treatment", "physiological", "physician"]


#helper method to ask user to enter directory to scan and check whether
# this exist or is valid
def helper_path():
    attempts = 0
    while attempts < 5:
        user_input = input(r'Enter directory to scan, for example C:\Windows\Documents: ')
        scan_path = Path(user_input)  #conver user input string to a path object

        # check if path exist and is a folder
        if scan_path.exists() and scan_path.is_dir():
            return scan_path  # Return as Path (safer than raw string)
        else:
            if attempts < 4:
                print('Invalid directory, check and try again')
                attempts += 1
            else:
                print('Invalid directory, try again later, goodbye!')
                sleep(3)
                exit()


#inform user to close any open application. If they want to proceed, ask for
#directory to scan. This function returns the path
def start():
    print('''
Welcome to GdprSafe v1.11 by EmilioB @All rights reserved
This application is resource intensive, if you have lot of data to scan I recommend you close
any application in use before initiate scanning.
Also make sure you have access to the folder you want to scan.
''')
    user_choice = input ("Do you want to continue? y/n (press Enter for y): ")

    if user_choice == "y" or user_choice == "":# if user enter y or press Enter
        return helper_path()#call helperPath() to verify directory exist
    elif user_choice == "n":# if user enter n, exit after 3 seconds
        print("Goodbye!")
        sleep(3)
        exit()
    else:
        # if user type any other char, give a second chance
        user_choice = input("Invalid entry. Type y/n or Enter for y :")
        if user_choice == "y" or user_choice == "":
            return helper_path()
        else:
            #if again a invalid entry, exit
            print("Invalid entry! Try again later, goodbye.")
            sleep(3)
            exit()


#check the directory for file with specific extensions and add them to the relevant lists.
# Print a summary of the number of files by type
def file_type():
    txt_count = 0
    csv_count = 0
    pdf_count = 0
    docx_count = 0
    excel_count = 0

    path_object = start()

    # use rglob('*') to recursively lists everything under path_object
    for p in path_object.rglob('*'):
        if not p.is_file(): #only files, skip directories
            continue

        file_name = p.name.lower()
        if fnmatch.fnmatch(file_name, '*.txt'):
            txt_files.append(p)
            txt_count += 1
        elif fnmatch.fnmatch(file_name, '*.csv'):
            csv_files.append(p)
            csv_count += 1
        elif fnmatch.fnmatch(file_name, '*.pdf'):
            pdf_files.append(p)
            pdf_count += 1
        elif fnmatch.fnmatch(file_name, '*.xlsx'):
            excel_files.append(p)
            excel_count += 1
        elif fnmatch.fnmatch(file_name, '*.docx'):
            docx_files.append(p)
            docx_count += 1

    print('Summary of files found by extension:')
    print(f'txt: {txt_count}')
    print(f'csv: {csv_count}')
    print(f'pdf: {pdf_count}')
    print(f'docx: {docx_count}')
    print(f'excel: {excel_count}')


#parse txt file list seeking key words
def txt_parse():
    txt_key_words = []
    list_of_lists = []
    print('Reading txt files...')

    for txt in tqdm(txt_files):
        # read_text() is the pathlib way to read file contents
        #https://stackoverflow.com/questions/78987882/is-pathlib-path-read-text-better-than-io-textiowrapper-read
        # errors="ignore" prevents UnicodeDecode errors from stopping the scan
        #https://stackoverflow.com/questions/24616678/unicodedecodeerror-in-python-when-reading-a-file-how-to-ignore-the-error-and-ju
        file_text = txt.read_text(errors="ignore")

        #keep path (txt) as string in the report for compatibility
        txt_key_words.append(str(txt))

        lower_text = file_text.lower()

        for word in words:
            if (word in lower_text) and (word not in txt_key_words):
                txt_key_words.append(word)

        list_of_lists.append(txt_key_words)
        txt_key_words = []  #reset for next file
        sleep(1)
    return list_of_lists


#parse pdf file list seeking key words
def pdf_parse():
    pdf_key_words = []
    list_of_lists = []
    print("Reading pdf files...")
    for pdf in tqdm(pdf_files):
        pdf_key_words.append(pdf) #add file path to the list
        document = pymupdf.open(pdf)
        for page in document:
            text = page.get_text().split()#list of words
            for word in text:
                word_lowercase = word.lower()
                if word_lowercase in words and word_lowercase not in pdf_key_words:
                    pdf_key_words.append(word_lowercase)
            list_of_lists.append(pdf_key_words)
            pdf_key_words = []
    sleep(1)
    return list_of_lists


#parse excel file list seeking key words
def excel_parse():
    xlsx_key_words = []
    list_of_lists = []
    print("Reading excel files...")
    for xlsx in tqdm(excel_files):
        wb = openpyxl.load_workbook(xlsx) #workbook
        xlsx_key_words.append(xlsx)
        ws = wb.worksheets #return a list of sheets in the wb
        for sheet in ws:
            for row in sheet:
                for cell in row:
                    for word in words:
                        if word == cell.value and word not in xlsx_key_words:
                            xlsx_key_words.append(word) #add key word to the list
        list_of_lists.append(xlsx_key_words)
        xlsx_key_words = []
    sleep(1)
    return list_of_lists

#parse csv file list seeking key words
def csv_parse():
    csv_key_words = []
    list_of_lists = []
    print('Reading csv files...')

    for cs in tqdm(csv_files):
        csv_key_words.append(str(cs))

        # path.open() is the pathlib equivalent of open()
        with cs.open('r', encoding='utf-8', newline='', errors='ignore') as f:
            reader = csv.reader(f)
            for line in reader:
                #cells to lowercase to match words consistently
                #use list comprehension
                lower_cells = [cell.lower() for cell in line]
                for word in words:
                    if word in lower_cells and word not in csv_key_words:
                        csv_key_words.append(word)

        list_of_lists.append(csv_key_words)
        csv_key_words = []
        sleep(1)

    return list_of_lists

#parse docx file list seeking key words
def docx_parse():
    print("Reading word files...")
    docx_key_words = []
    list_of_lists = []
    for doc in tqdm(docx_files):
        doc_open = docx.Document(doc)
        docx_key_words.append(doc)

        for par in doc_open.paragraphs: #paragraphs
            for p in par.runs:
                p_text = p.text.lower()
                for word in words:
                    if (word in p_text) and (word not in docx_key_words):
                        docx_key_words.append(word)
        list_of_lists.append(docx_key_words)
        docx_key_words = []
    sleep(1)
    return list_of_lists

#add each file link and their keywords in a single line in csv file
# parameter: list of lists returned by each scanner
def add_to_report(_key_words):
    print("Writing to csv report...")

    report_path = Path("report.csv")#Path object for the report file

    #use newline="" for csv on windows to avoid blank lines
    with report_path.open('a', newline='', encoding='utf-8') as report:
        csv_writer = csv.writer(report)
        for w in _key_words:
            csv_writer.writerow(w)

    print('Complete!')

#final function
def end():
    work_dir = Path.cwd()#current working directory as a Path
    print(f'''
Pick your csv report in the {work_dir} folder. Goodbye!
''')
    input("Close the window or press enter to exit the program")

#main
def main():
    file_type()
    print("\n")

    txt_key_words = txt_parse()
    add_to_report(txt_key_words)
    print("\n")

    pdf_key_words = pdf_parse()
    add_to_report(pdf_key_words)
    print("\n")

    xlsx_key_words = excel_parse()
    add_to_report(xlsx_key_words)
    print()

    csv_key_words = csv_parse()
    add_to_report(csv_key_words)
    print("\n")

    docx_key_words = docx_parse()
    add_to_report(docx_key_words)
    end()
         

if __name__ == "__main__":
    main()

